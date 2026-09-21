from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .extractor import EXPECTED_HEADERS, CauseRow


WIDTHS = [9, 22, 42, 7, 10, 9, 32, 42]


def write_excel(rows: list[CauseRow], target_date: date, output: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cause List"

    ws.merge_cells("A1:H1")
    ws["A1"] = target_date.strftime("%d.%m.%Y")
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True, underline="single")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    for col, header in enumerate(EXPECTED_HEADERS, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = Font(name="Times New Roman", size=10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin = Side(style="thin", color="000000")
    for row in ws.iter_rows(min_row=3, max_row=3, min_col=1, max_col=8):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for r_idx, cause_row in enumerate(rows, start=4):
        for c_idx, value in enumerate(cause_row.values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        # Height approximation based on the amount of wrapped text.
        max_lines = max(1, max((len(v) // max(1, WIDTHS[i]) + 1) for i, v in enumerate(cause_row.values)))
        ws.row_dimensions[r_idx].height = min(120, max(18, 15 * max_lines))

    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
